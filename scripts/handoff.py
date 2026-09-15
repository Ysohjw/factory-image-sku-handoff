#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Factory Image SKU Handoff contributors
"""Bounded, offline XLSX image-to-declared-row evidence exporter.

Original implementation. Workbook semantics and drawing models are parsed by
openpyxl; lxml/ZIP inspection guards the package and reconciles its inventory.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, time
from decimal import Decimal
import hashlib
from io import BytesIO
import json
import math
from pathlib import Path, PurePosixPath
import posixpath
import re
import shutil
import stat
import sys
import tempfile
import warnings
import zipfile

from lxml import etree
import openpyxl
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.packaging.relationship import RelationshipList
from openpyxl.packaging.workbook import WorkbookPackage
from openpyxl.utils.cell import column_index_from_string, coordinate_to_tuple, get_column_letter, range_boundaries
from PIL import Image

VERSION = "0.1.1-beta.1"
MIB = 1024 * 1024
LIMITS = {
    "archive_bytes": 20 * MIB, "members": 2000, "expanded_bytes": 100 * MIB,
    "member_bytes": 20 * MIB, "xml_bytes": 5 * MIB, "xml_nodes_part": 100000,
    "xml_nodes_total": 500000, "cells": 100000, "merged_area": 100000,
    "sheets": 50, "rows": 10000, "columns": 256, "media": 1000,
    "occurrences": 2000, "image_bytes": 10 * MIB, "image_pixels": 25000000,
    "sku_characters": 4096,
}
S = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
D = "{http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing}"
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
Image.MAX_IMAGE_PIXELS = LIMITS["image_pixels"]


class Rejected(Exception):
    def __init__(self, code):
        self.code = code
        super().__init__(code)


def require(condition, code):
    if not condition:
        raise Rejected(code)


def sha(data):
    return hashlib.sha256(data).hexdigest()


def json_bytes(value):
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + "\n").encode("utf-8")


def local_name(tag):
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def valid_member(name, directory=False):
    require(isinstance(name, str) and bool(name), "UNSAFE_MEMBER_NAME")
    require(not name.startswith(("/", "\\")) and "\\" not in name and ":" not in name,
            "UNSAFE_MEMBER_NAME")
    require(not any(ord(ch) < 32 or ord(ch) == 127 for ch in name), "UNSAFE_MEMBER_NAME")
    value = name[:-1] if directory and name.endswith("/") else name
    require(all(p not in ("", ".", "..") for p in value.split("/")), "UNSAFE_MEMBER_NAME")
    return value


def xml_parse(data):
    require(len(data) <= LIMITS["xml_bytes"], "XML_BYTES_LIMIT")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise Rejected("XML_ENCODING_UNSUPPORTED")
    require("\x00" not in text, "XML_ENCODING_UNSUPPORTED")
    declaration = re.match(r"\s*<\?xml\s+([^?]*)\?>", text)
    if declaration:
        encoding = re.search(r"encoding\s*=\s*['\"]([^'\"]+)['\"]", declaration.group(1), re.I)
        if encoding:
            require(encoding.group(1).lower() in ("utf-8", "utf8", "ascii", "us-ascii"),
                    "XML_ENCODING_UNSUPPORTED")
    require(not re.search(r"<!\s*(?:DOCTYPE|ENTITY)\b", text, re.I), "XML_DTD_ENTITY_FORBIDDEN")
    parser = etree.XMLParser(resolve_entities=False, no_network=True, load_dtd=False,
                             huge_tree=False, recover=False)
    try:
        root = etree.fromstring(data, parser=parser)
    except (etree.XMLSyntaxError, ValueError):
        raise Rejected("INVALID_XML")
    count = sum(1 for e in root.iter() if isinstance(e.tag, str))
    require(count <= LIMITS["xml_nodes_part"], "XML_NODES_PART_LIMIT")
    return root, count


def part_for_rels(name):
    if name == "_rels/.rels":
        return ""
    p = PurePosixPath(name)
    require(p.parent.name == "_rels" and p.name.endswith(".rels"), "INVALID_RELATIONSHIP_PART")
    return str(p.parent.parent / p.name[:-5])


def resolve_target(source, target):
    require(bool(target) and "\\" not in target and not any(ord(c) < 32 for c in target),
            "INVALID_RELATIONSHIP_TARGET")
    require(":" not in target and "?" not in target and "#" not in target,
            "EXTERNAL_OR_UNSUPPORTED_RELATIONSHIP")
    value = target[1:] if target.startswith("/") else posixpath.join(posixpath.dirname(source), target)
    value = posixpath.normpath(value)
    valid_member(value)
    return value


def bounds(ref):
    try:
        a, b, c, d = range_boundaries(ref)
    except (ValueError, TypeError):
        raise Rejected("INVALID_WORKSHEET_RANGE")
    require(all(isinstance(x, int) for x in (a, b, c, d)), "INVALID_WORKSHEET_RANGE")
    require(1 <= a <= c <= LIMITS["columns"] and 1 <= b <= d <= LIMITS["rows"],
            "WORKSHEET_DIMENSION_LIMIT")
    return a, b, c, d


def package_guard(data):
    """Read ZIP entries in memory; never extract archive paths to disk."""
    try:
        archive = zipfile.ZipFile(BytesIO(data))
    except (zipfile.BadZipFile, ValueError):
        raise Rejected("INVALID_ZIP")
    parts, xml, rels = {}, {}, {}
    expanded = nodes = cells = merged = worksheets = 0
    infos = archive.infolist()
    require(len(infos) <= LIMITS["members"], "MEMBER_COUNT_LIMIT")
    seen = set()
    for info in infos:
        name = info.filename
        require(info.orig_filename == name, "UNSAFE_MEMBER_NAME")
        canonical = valid_member(name, info.is_dir())
        require(canonical.casefold() not in seen, "DUPLICATE_MEMBER")
        seen.add(canonical.casefold())
        require(not info.flag_bits & 1, "ENCRYPTED_ZIP_FORBIDDEN")
        mode = stat.S_IFMT(info.external_attr >> 16)
        require(mode in (0, stat.S_IFREG, stat.S_IFDIR), "NONREGULAR_MEMBER_FORBIDDEN")
        require(mode != stat.S_IFDIR or info.is_dir(), "NONREGULAR_MEMBER_FORBIDDEN")
        require(info.compress_type in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED), "ZIP_COMPRESSION_UNSUPPORTED")
        require(0 <= info.file_size <= LIMITS["member_bytes"], "MEMBER_BYTES_LIMIT")
        require(info.file_size == 0 or info.compress_size > 0, "INVALID_COMPRESSED_SIZE")
        require(info.file_size <= MIB or info.file_size <= 200 * info.compress_size, "COMPRESSION_RATIO_LIMIT")
        expanded += info.file_size
        require(expanded <= LIMITS["expanded_bytes"], "EXPANDED_BYTES_LIMIT")
        lower = name.lower()
        require(not any(v in lower for v in ("vbaproject", "/activex/", "/embeddings/"))
                and not lower.endswith((".exe", ".dll", ".com", ".scr", ".js", ".vbs", ".ps1")),
                "ACTIVE_CONTENT_FORBIDDEN")
        if info.is_dir():
            continue
        try:
            with archive.open(info) as handle:
                payload = handle.read(LIMITS["member_bytes"] + 1)
        except (zipfile.BadZipFile, RuntimeError, OSError, EOFError):
            raise Rejected("INVALID_ZIP_MEMBER")
        require(len(payload) == info.file_size and len(payload) <= LIMITS["member_bytes"], "MEMBER_SIZE_MISMATCH")
        parts[name] = payload
        if lower.endswith((".xml", ".rels", ".vml")):
            root, count = xml_parse(payload)
            xml[name] = root
            nodes += count
            require(nodes <= LIMITS["xml_nodes_total"], "XML_NODES_TOTAL_LIMIT")
            if root.tag == S + "worksheet":
                worksheets += 1
                require(worksheets <= LIMITS["sheets"], "SHEET_COUNT_LIMIT")
                seen_rows, seen_cells = set(), set()
                for e in root.iter():
                    if e.tag == S + "dimension":
                        bounds(e.get("ref", ""))
                    elif e.tag == S + "row":
                        try:
                            row = int(e.get("r", "0"))
                        except ValueError:
                            raise Rejected("INVALID_ROW")
                        require(1 <= row <= LIMITS["rows"], "WORKSHEET_DIMENSION_LIMIT")
                        require(row not in seen_rows, "DUPLICATE_ROW_INDEX")
                        seen_rows.add(row)
                    elif e.tag == S + "c":
                        x1, y1, x2, y2 = bounds(e.get("r", ""))
                        require(x1 == x2 and y1 == y2, "INVALID_CELL_COORDINATE")
                        coordinate = (y1, x1)
                        require(coordinate not in seen_cells, "DUPLICATE_CELL_COORDINATE")
                        seen_cells.add(coordinate)
                        parent = e.getparent()
                        require(parent is not None and parent.tag == S + "row", "CELL_ROW_MISMATCH")
                        try:
                            parent_row = int(parent.get("r", "0"))
                        except ValueError:
                            raise Rejected("CELL_ROW_MISMATCH")
                        require(y1 == parent_row, "CELL_ROW_MISMATCH")
                        cells += 1
                        require(cells <= LIMITS["cells"], "WORKSHEET_CELL_LIMIT")
                    elif e.tag == S + "mergeCell":
                        x1, y1, x2, y2 = bounds(e.get("ref", ""))
                        merged += (x2 - x1 + 1) * (y2 - y1 + 1)
                        require(merged <= LIMITS["merged_area"], "MERGED_AREA_LIMIT")
                    elif e.tag == S + "col":
                        try:
                            start, end = int(e.get("min", "0")), int(e.get("max", "0"))
                        except ValueError:
                            raise Rejected("INVALID_COLUMN")
                        require(1 <= start <= end <= LIMITS["columns"], "WORKSHEET_DIMENSION_LIMIT")
                    for attr in ("ht", "width", "defaultRowHeight", "defaultColWidth", "baseColWidth"):
                        if attr in e.attrib:
                            try:
                                number = float(e.attrib[attr])
                            except ValueError:
                                raise Rejected("INVALID_GEOMETRY")
                            require(math.isfinite(number) and number >= 0, "INVALID_GEOMETRY")
    archive.close()
    require("[Content_Types].xml" in xml and "_rels/.rels" in xml, "MISSING_PACKAGE_ROOTS")
    ct = xml["[Content_Types].xml"]
    xml_defaults = set()
    for e in ct:
        kind = e.get("ContentType", "").lower()
        require(not any(s in kind for s in ("macroenabled", "vbaproject", "activex", "oleobject")),
                "ACTIVE_CONTENT_FORBIDDEN")
        # OpenXML package part names need not end in .xml. Check content-type
        # declarations too so renamed styles/drawings cannot bypass the guard.
        if kind.endswith("xml"):
            if e.get("Extension"):
                xml_defaults.add("." + e.get("Extension").lower())
            if e.get("PartName"):
                member = e.get("PartName").lstrip("/")
                require(member in parts, "BROKEN_CONTENT_TYPE_PART")
                if member not in xml:
                    root, count = xml_parse(parts[member])
                    # Workbook/sheet roots are supported only after the full
                    # worksheet-bound guard above has inspected them.
                    require(root.tag not in (S + "worksheet", S + "workbook"), "NONSTANDARD_WORKBOOK_PART_UNSUPPORTED")
                    xml[member] = root
                    nodes += count
                    require(nodes <= LIMITS["xml_nodes_total"], "XML_NODES_TOTAL_LIMIT")
    for member in parts:
        if member not in xml and any(member.lower().endswith(ext) for ext in xml_defaults):
            root, count = xml_parse(parts[member])
            require(root.tag not in (S + "worksheet", S + "workbook"), "NONSTANDARD_WORKBOOK_PART_UNSUPPORTED")
            xml[member] = root
            nodes += count
            require(nodes <= LIMITS["xml_nodes_total"], "XML_NODES_TOTAL_LIMIT")
    for name in sorted(xml):
        if not name.lower().endswith(".rels"):
            continue
        source = part_for_rels(name)
        require(not source or source in parts, "ORPHAN_RELATIONSHIP_PART")
        root = xml[name]
        require(root.tag == "{" + REL_NS + "}Relationships", "INVALID_RELATIONSHIP_PART")
        raw = list(root)
        try:
            parsed = RelationshipList.from_tree(root)
        except (TypeError, ValueError, AttributeError):
            raise Rejected("INVALID_RELATIONSHIP")
        require(len(parsed) == len(raw), "PARSER_INVENTORY_MISMATCH")
        records = {}
        for relationship in parsed:
            require(relationship.Id and relationship.Id not in records, "DUPLICATE_RELATIONSHIP")
            require(relationship.TargetMode in (None, "Internal"), "EXTERNAL_RELATIONSHIP_FORBIDDEN")
            target = resolve_target(source, relationship.Target)
            require(target in parts, "BROKEN_INTERNAL_RELATIONSHIP")
            require(target in xml or relationship.Type.endswith(("/image", "/printerSettings")),
                    "UNSUPPORTED_OPAQUE_RELATIONSHIP")
            records[relationship.Id] = {"target": target, "type": relationship.Type}
        rels[source] = records
    roots = [r["target"] for r in rels.get("", {}).values() if r["type"].endswith("/officeDocument")]
    require(len(roots) == 1 and roots[0] in xml, "MISSING_WORKBOOK_ROOT")
    require(xml[roots[0]].tag == S + "workbook", "UNSUPPORTED_WORKBOOK_NAMESPACE")
    return parts, xml, rels, roots[0]


def asset_inventory(parts, xml, rels):
    media = {p for p in parts if p.lower().startswith("xl/media/")}
    for links in rels.values():
        media.update(r["target"] for r in links.values() if r["type"].endswith("/image"))
    content = xml["[Content_Types].xml"]
    for entry in content:
        if entry.get("ContentType", "").startswith("image/"):
            if entry.get("PartName"):
                media.add(entry.get("PartName").lstrip("/"))
            elif entry.get("Extension"):
                ext = "." + entry.get("Extension").lower()
                media.update(p for p in parts if p.lower().endswith(ext))
    require(len(media) <= LIMITS["media"], "MEDIA_COUNT_LIMIT")
    assets, by_member, payloads = {}, {}, {}
    for member in sorted(media):
        require(member in parts, "BROKEN_MEDIA_PART")
        raw = parts[member]
        digest = sha(raw)
        detected = "png" if raw.startswith(b"\x89PNG\r\n\x1a\n") else "jpg" if raw.startswith(b"\xff\xd8\xff") else None
        reasons = []
        if detected:
            require(len(raw) <= LIMITS["image_bytes"], "IMAGE_BYTES_LIMIT")
            try:
                with warnings.catch_warnings():
                    warnings.simplefilter("error", Image.DecompressionBombWarning)
                    with Image.open(BytesIO(raw)) as im:
                        require(im.format in ("PNG", "JPEG"), "IMAGE_FORMAT_MISMATCH")
                        require(im.width > 0 and im.height > 0 and im.width * im.height <= LIMITS["image_pixels"], "IMAGE_PIXELS_LIMIT")
                        im.verify()
            except Rejected:
                raise
            except (Image.DecompressionBombError, Image.DecompressionBombWarning):
                raise Rejected("IMAGE_PIXELS_LIMIT")
            except Exception:
                raise Rejected("INVALID_IMAGE")
            path = "assets/" + digest + "." + detected
            payloads[path] = raw
        else:
            path = None
            reasons.append("UNSUPPORTED_MEDIA")
        if digest not in assets:
            assets[digest] = {"asset_id": digest, "sha256": digest, "byte_length": len(raw),
                              "format": detected, "output_path": path, "package_members": [],
                              "reason_codes": reasons}
        assets[digest]["package_members"].append(member)
        by_member[member] = digest
    return assets, by_member, payloads


def sku_evidence(cell):
    value = cell.value
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    elif value is not None and not isinstance(value, (str, int, float, bool)):
        value = str(value)
    if isinstance(value, float):
        require(math.isfinite(value), "NONFINITE_CELL_VALUE")
    if isinstance(value, str):
        require(len(value) <= LIMITS["sku_characters"], "SKU_LENGTH_LIMIT")
    return {"value": value, "data_type": cell.data_type, "number_format": cell.number_format}


def sku_reasons(evidence):
    if evidence["value"] is None or evidence["value"] == "":
        return ["SKU_BLANK"]
    if evidence["data_type"] not in ("s", "inlineStr") or not isinstance(evidence["value"], str):
        return ["SKU_NOT_TEXT"]
    return []


def hidden_column(ws, col):
    for key, dim in ws.column_dimensions.items():
        start = dim.min or column_index_from_string(key)
        end = dim.max or start
        if start <= col <= end and dim.hidden:
            return True
    return False


def marker_dict(marker):
    return {"row": marker.row, "column": marker.col, "row_offset_emu": marker.rowOff,
            "column_offset_emu": marker.colOff}


def geometry(anchor, kind, ws, image_cols):
    evidence, reasons = {"kind": kind}, []
    if kind == "absoluteAnchor":
        evidence["position_emu"] = {"x": anchor.pos.x, "y": anchor.pos.y}
        evidence["extent_emu"] = {"width": anchor.ext.cx, "height": anchor.ext.cy}
        return None, None, evidence, ["ABSOLUTE_ANCHOR"]
    start = anchor._from
    evidence["from"] = marker_dict(start)
    row, col = start.row + 1, start.col + 1
    require(0 <= start.row < LIMITS["rows"] and 0 <= start.col < LIMITS["columns"], "ANCHOR_DIMENSION_LIMIT")
    if start.rowOff < 0 or start.colOff < 0:
        reasons.append("NEGATIVE_ANCHOR_OFFSET")
    if col not in image_cols:
        reasons.append("OUTSIDE_IMAGE_COLUMNS")
    if kind == "oneCellAnchor":
        evidence["extent_emu"] = {"width": anchor.ext.cx, "height": anchor.ext.cy}
        dimension = ws.row_dimensions.get(row) if ws is not None else None
        height = dimension.height if dimension is not None else None
        evidence["explicit_row_height_points"] = height
        evidence["horizontal_rule"] = "declared_start_column_only"
        if anchor.ext.cx <= 0 or anchor.ext.cy <= 0:
            reasons.append("NONPOSITIVE_IMAGE_EXTENT")
        if height is None or not math.isfinite(height) or height <= 0:
            reasons.append("ROW_HEIGHT_UNKNOWN")
        elif Decimal(start.rowOff + anchor.ext.cy) > Decimal(str(height)) * 12700:
            reasons.append("IMAGE_SPANS_ROWS")
    else:
        end = anchor.to
        evidence["to"] = marker_dict(end)
        require(0 <= end.row <= LIMITS["rows"] and 0 <= end.col <= LIMITS["columns"], "ANCHOR_DIMENSION_LIMIT")
        if end.rowOff < 0 or end.colOff < 0:
            reasons.append("NEGATIVE_ANCHOR_OFFSET")
        row_inside = end.row == start.row and end.rowOff > start.rowOff or end.row == start.row + 1 and end.rowOff == 0
        col_inside = end.col == start.col and end.colOff > start.colOff or end.col == start.col + 1 and end.colOff == 0
        if not row_inside:
            reasons.append("IMAGE_SPANS_ROWS_OR_INVALID_EXTENT")
        if not col_inside:
            reasons.append("IMAGE_SPANS_COLUMNS_OR_INVALID_EXTENT")
        dimension = ws.row_dimensions.get(row) if ws is not None else None
        height = dimension.height if dimension is not None else None
        evidence["explicit_row_height_points"] = height
        if start.rowOff or end.row == start.row and end.rowOff:
            if height is None or not math.isfinite(height) or height <= 0:
                reasons.append("ROW_HEIGHT_UNKNOWN")
            elif Decimal(start.rowOff) >= Decimal(str(height)) * 12700 or end.row == start.row and Decimal(end.rowOff) > Decimal(str(height)) * 12700:
                reasons.append("IMAGE_SPANS_ROWS")
        if start.colOff or end.colOff:
            reasons.append("TWO_CELL_HORIZONTAL_OFFSET_REVIEW")
    return row, col, evidence, reasons


def analyze(data, layout):
    parts, xml, rels, workbook_part = package_guard(data)
    assets, by_member, payloads = asset_inventory(parts, xml, rels)
    issues = []

    def issue(code, scope):
        entry = {"code": code, "scope": scope}
        if entry not in issues:
            issues.append(entry)

    try:
        package = WorkbookPackage.from_tree(xml[workbook_part])
    except (TypeError, ValueError, AttributeError):
        raise Rejected("INVALID_WORKBOOK")
    require(len(package.sheets) <= LIMITS["sheets"], "SHEET_COUNT_LIMIT")
    sheet_parts = {}
    for sheet in package.sheets:
        link = rels.get(workbook_part, {}).get(sheet.id)
        require(link is not None and sheet.name not in sheet_parts, "INVALID_SHEET_RELATIONSHIP")
        sheet_parts[sheet.name] = link["target"]
    require(layout["sheet"] in sheet_parts, "SHEET_NOT_FOUND")
    try:
        with warnings.catch_warnings(record=True) as notices:
            warnings.simplefilter("always")
            wb = openpyxl.load_workbook(BytesIO(data), read_only=False, data_only=False, keep_links=False, rich_text=False)
        if notices:
            issue("PARSER_WARNING", "workbook")
    except Exception:
        raise Rejected("WORKBOOK_PARSER_REJECTED")
    require(layout["sheet"] in wb.sheetnames, "SHEET_NOT_FOUND")
    selected = wb[layout["sheet"]]
    require(hasattr(selected, "_cells"), "SELECTED_SHEET_UNSUPPORTED")
    image_cols = [column_index_from_string(v) for v in layout["image_columns"]]
    sku_col = column_index_from_string(layout["sku_column"])
    first = layout["first_data_row"]

    owners = defaultdict(list)
    for sheet_name, part in sheet_parts.items():
        if part not in xml or xml[part].tag != S + "worksheet":
            issue("UNSUPPORTED_SHEET_KIND", sheet_name)
            continue
        for e in xml[part].iter():
            if e.tag == S + "drawing":
                rid = e.get(R + "id")
                link = rels.get(part, {}).get(rid)
                require(link is not None and link["type"].endswith("/drawing"), "BROKEN_DRAWING_RELATIONSHIP")
                owners[link["target"]].append(sheet_name)
            if e.tag == S + "f" and re.search(r"(?:DISPIMG|(?:_xlfn\.)?IMAGE)\s*\(", e.text or "", re.I):
                issue("UNSUPPORTED_IN_CELL_IMAGE_FORMULA", sheet_name)
            if e.tag == S + "c" and (e.get("vm") is not None or e.get("cm") is not None):
                issue("UNSUPPORTED_CELL_METADATA", sheet_name)
    for name in parts:
        lower = name.lower()
        if "richdata" in lower or "cellimages" in lower or lower.endswith(".vml"):
            issue("UNSUPPORTED_IMAGE_LAYOUT", name)

    occurrences, used_members, consumed_refs = [], set(), set()

    def add_occurrence(part, rid, sheet_name, anchor=None, kind=None, extra=None):
        link = rels.get(part, {}).get(rid)
        require(link is not None and link["type"].endswith("/image"), "BROKEN_IMAGE_RELATIONSHIP")
        member = link["target"]
        require(member in by_member, "PARSER_INVENTORY_MISMATCH")
        used_members.add(member)
        consumed_refs.add((part, rid))
        reasons = list(extra or [])
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else None
        row = col = None
        evidence = {"kind": "unsupported_or_unanchored"}
        if anchor is not None:
            row, col, evidence, gr = geometry(anchor, kind, ws, image_cols)
            reasons.extend(gr)
        else:
            reasons.append("UNSUPPORTED_OR_UNANCHORED_IMAGE")
        if sheet_name != layout["sheet"]:
            reasons.append("OUTSIDE_SELECTED_SHEET" if sheet_name else "ORPHAN_DRAWING")
        if row is not None and row < first:
            reasons.append("BEFORE_FIRST_DATA_ROW")
        digest = by_member[member]
        reasons.extend(assets[digest]["reason_codes"])
        evidence_sku = None
        cell_name = None
        if row is not None and ws is not None and hasattr(ws, "cell"):
            cell_name = get_column_letter(sku_col) + str(row)
            cell = ws.cell(row, sku_col)
            evidence_sku = sku_evidence(cell)
            reasons.extend(sku_reasons(evidence_sku))
            if ws.sheet_state != "visible":
                reasons.append("HIDDEN_SHEET")
            dimension = ws.row_dimensions.get(row)
            if dimension is not None and dimension.hidden:
                reasons.append("HIDDEN_ROW")
            if hidden_column(ws, sku_col) or col is not None and hidden_column(ws, col):
                reasons.append("HIDDEN_COLUMN")
            for merged in ws.merged_cells.ranges:
                if merged.min_row <= row <= merged.max_row and any(merged.min_col <= c <= merged.max_col for c in (sku_col, col) if c is not None):
                    reasons.append("MERGED_CANDIDATE_CELL")
                    break
        occurrences.append({"occurrence_id": f"image-{len(occurrences) + 1:04d}", "source_part": part,
                            "relationship_id": rid, "sheet": sheet_name, "asset_id": digest,
                            "anchor": evidence, "candidate_row": row, "sku_cell": cell_name,
                            "sku": evidence_sku, "disposition": "review" if reasons else "eligible",
                            "reason_codes": sorted(set(reasons))})
        require(len(occurrences) <= LIMITS["occurrences"], "OCCURRENCE_COUNT_LIMIT")

    for part in sorted(xml):
        tree = xml[part]
        raw_blips = tree.findall(".//" + A + "blip")
        if tree.tag != D + "wsDr":
            for blip in raw_blips:
                require(blip.get(R + "link") is None, "EXTERNAL_IMAGE_FORBIDDEN")
                rid = blip.get(R + "embed")
                if rid:
                    add_occurrence(part, rid, None, extra=["UNSUPPORTED_IMAGE_LAYOUT"])
            continue
        try:
            drawing = SpreadsheetDrawing.from_tree(tree)
        except (TypeError, ValueError, AttributeError):
            raise Rejected("DRAWING_PARSER_REJECTED")
        models = {"twoCellAnchor": drawing.twoCellAnchor, "oneCellAnchor": drawing.oneCellAnchor,
                  "absoluteAnchor": drawing.absoluteAnchor}
        indices = Counter()
        represented = []
        for node in tree:
            kind = local_name(node.tag)
            if kind not in models:
                if isinstance(node.tag, str):
                    issue("UNSUPPORTED_DRAWING_OBJECT", part)
                continue
            index = indices[kind]
            indices[kind] += 1
            require(index < len(models[kind]), "PARSER_INVENTORY_MISMATCH")
            anchor = models[kind][index]
            pics = node.findall(".//" + A + "blip")
            model_pic = anchor.pic or (anchor.groupShape.pic if anchor.groupShape is not None else None)
            model_blip = model_pic.blipFill.blip if model_pic is not None and model_pic.blipFill is not None else None
            require(len(pics) == (1 if model_blip is not None else 0), "PARSER_INVENTORY_MISMATCH")
            if not pics:
                issue("UNSUPPORTED_DRAWING_OBJECT", part)
                continue
            blip = pics[0]
            require(blip.get(R + "link") is None, "EXTERNAL_IMAGE_FORBIDDEN")
            rid = blip.get(R + "embed")
            require(rid is not None and model_blip.embed == rid, "PARSER_INVENTORY_MISMATCH")
            represented.append(rid)
            extra = ["GROUPED_DRAWING_REVIEW"] if anchor.groupShape is not None else []
            if any(local_name(child.tag) not in ("from", "to", "pos", "ext", "pic", "clientData") for child in node if isinstance(child.tag, str)):
                extra.append("UNSUPPORTED_DRAWING_OBJECT")
            for sheet_name in owners.get(part, [None]):
                add_occurrence(part, rid, sheet_name, anchor, kind, extra)
        require(Counter(b.get(R + "embed") for b in raw_blips) == Counter(represented), "PARSER_INVENTORY_MISMATCH")
    for part in sorted(rels):
        for rid, link in sorted(rels[part].items()):
            if link["type"].endswith("/image") and (part, rid) not in consumed_refs:
                add_occurrence(part, rid, None, extra=["UNSUPPORTED_IMAGE_REFERENCE"])
    for member, digest in by_member.items():
        if member not in used_members:
            assets[digest]["reason_codes"] = sorted(set(assets[digest]["reason_codes"] + ["UNREFERENCED_MEDIA"]))
            issue("UNREFERENCED_MEDIA", member)
        if assets[digest]["format"] is None:
            issue("UNSUPPORTED_MEDIA", member)
    if not occurrences:
        issue("NO_IMAGE_OCCURRENCES", "workbook")
    if not any(a["format"] for a in assets.values()):
        issue("NO_SUPPORTED_IMAGES", "workbook")

    selected_occ = defaultdict(list)
    for occurrence in occurrences:
        if occurrence["sheet"] == layout["sheet"] and occurrence["candidate_row"] is not None:
            selected_occ[occurrence["candidate_row"]].append(occurrence)
    candidate_rows = {row for (row, col), cell in selected._cells.items() if col == sku_col and row >= first and cell.value is not None}
    candidate_rows.update(row for row in selected_occ if row >= first)
    rows = []
    texts = Counter()
    for row in sorted(candidate_rows):
        evidence = sku_evidence(selected.cell(row, sku_col))
        if not sku_reasons(evidence):
            texts[evidence["value"]] += 1
        rows.append({"row": row, "sku_cell": get_column_letter(sku_col) + str(row), "sku": evidence,
                     "image_count": len(selected_occ[row]), "disposition": "eligible", "reason_codes": []})
    for record in rows:
        reasons = sku_reasons(record["sku"])
        row = record["row"]
        if record["image_count"] == 0:
            reasons.append("SKU_ROW_WITHOUT_IMAGE")
        elif record["image_count"] > 1:
            reasons.append("MULTIPLE_IMAGES_SAME_ROW")
        if not sku_reasons(record["sku"]) and texts[record["sku"]["value"]] > 1:
            reasons.append("DUPLICATE_SKU")
        if selected.sheet_state != "visible":
            reasons.append("HIDDEN_SHEET")
        dimension = selected.row_dimensions.get(row)
        if dimension is not None and dimension.hidden:
            reasons.append("HIDDEN_ROW")
        if hidden_column(selected, sku_col):
            reasons.append("HIDDEN_COLUMN")
        for merged in selected.merged_cells.ranges:
            if merged.min_row <= row <= merged.max_row and merged.min_col <= sku_col <= merged.max_col:
                reasons.append("MERGED_CANDIDATE_CELL")
                break
        for occurrence in selected_occ[row]:
            occurrence["reason_codes"] = sorted(set(occurrence["reason_codes"] + reasons))
            occurrence["disposition"] = "review" if occurrence["reason_codes"] else "eligible"
            reasons.extend(occurrence["reason_codes"])
        record["reason_codes"] = sorted(set(reasons))
        record["disposition"] = "review" if reasons else "eligible"
    wb.close()
    status = "review" if issues or any(o["reason_codes"] for o in occurrences) or any(r["reason_codes"] for r in rows) else "eligible"
    manifest = {"schema_version": "1.0", "tool_version": VERSION, "source_sha256": sha(data), "layout": layout,
                "status": status, "assets": [assets[k] for k in sorted(assets)], "occurrences": occurrences,
                "rows": rows, "issues": sorted(issues, key=lambda i: (i["code"], i["scope"]))}
    return manifest, payloads


class Parser(argparse.ArgumentParser):
    def error(self, message):
        raise Rejected("INVALID_ARGUMENTS")


def arguments(argv=None):
    parser = Parser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--sku-column", required=True)
    parser.add_argument("--image-columns", required=True)
    parser.add_argument("--first-data-row", type=int, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    image_cols = args.image_columns.split(",")
    require(bool(args.sheet) and 1 <= args.first_data_row <= LIMITS["rows"], "INVALID_ARGUMENTS")
    require(all(re.fullmatch(r"[A-Z]{1,3}", c) for c in [args.sku_column] + image_cols), "INVALID_ARGUMENTS")
    require(len(set(image_cols)) == len(image_cols) and args.sku_column not in image_cols, "INVALID_ARGUMENTS")
    require(all(column_index_from_string(c) <= LIMITS["columns"] for c in [args.sku_column] + image_cols), "INVALID_ARGUMENTS")
    require(args.input.suffix.lower() == ".xlsx" and args.input.is_file(), "INPUT_NOT_XLSX_FILE")
    require(not args.output.exists() and not args.output.is_symlink(), "OUTPUT_ALREADY_EXISTS")
    require(args.output.parent.is_dir(), "OUTPUT_PARENT_MISSING")
    source, output = args.input.resolve(), args.output.resolve()
    require(source != output and not source.is_relative_to(output) and not output.is_relative_to(source), "INPUT_OUTPUT_OVERLAP")
    layout = {"sheet": args.sheet, "sku_column": args.sku_column,
              "image_columns": sorted(image_cols, key=column_index_from_string), "first_data_row": args.first_data_row,
              "ownership_rule": "user_declared_row_layout_requires_human_semantic_confirmation"}
    return args, layout


def publish(output, manifest, payloads):
    require(not output.exists() and not output.is_symlink(), "OUTPUT_ALREADY_EXISTS")
    staging = Path(tempfile.mkdtemp(prefix=".factory-handoff-", dir=output.parent)).resolve()
    parent = output.parent.resolve()
    require(staging.parent == parent, "UNSAFE_STAGING_PATH")
    try:
        (staging / "assets").mkdir()
        for relative, data in sorted(payloads.items()):
            require(re.fullmatch(r"assets/[0-9a-f]{64}\.(png|jpg)", relative) is not None, "UNSAFE_OUTPUT_PATH")
            (staging / relative).write_bytes(data)
        (staging / "manifest.json").write_bytes(json_bytes(manifest))
        require(not output.exists() and not output.is_symlink(), "OUTPUT_ALREADY_EXISTS")
        # On the supported Windows runtime rename refuses an existing target.
        staging.rename(output)
    finally:
        if staging.exists():
            require(staging.parent == parent and staging.name.startswith(".factory-handoff-"), "UNSAFE_STAGING_PATH")
            shutil.rmtree(staging)


def main(argv=None):
    try:
        args, layout = arguments(argv)
        require(args.input.stat().st_size <= LIMITS["archive_bytes"], "ARCHIVE_BYTES_LIMIT")
        with args.input.open("rb") as handle:
            data = handle.read(LIMITS["archive_bytes"] + 1)
        require(len(data) <= LIMITS["archive_bytes"], "ARCHIVE_BYTES_LIMIT")
        manifest, payloads = analyze(data, layout)
        publish(args.output, manifest, payloads)
        print(json.dumps({"status": manifest["status"], "manifest": "manifest.json",
                          "assets": len(manifest["assets"]), "occurrences": len(manifest["occurrences"]),
                          "rows": len(manifest["rows"])}, sort_keys=True))
        return 0 if manifest["status"] == "eligible" else 2
    except Rejected as exc:
        print(json.dumps({"error": exc.code, "status": "rejected"}, sort_keys=True), file=sys.stderr)
        return 1
    except Exception:
        # Exception text can contain workbook strings and absolute paths.
        print(json.dumps({"error": "UNEXPECTED_INPUT_OR_IO_ERROR", "status": "rejected"}, sort_keys=True), file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
